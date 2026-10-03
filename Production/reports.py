import csv
import io
import secrets
from datetime import datetime

from flask import Blueprint, current_app, request, session, send_file, render_template, redirect, url_for, flash
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

from models import db_connection

reports_bp = Blueprint("reports", __name__, url_prefix="/reports")

MODULE_NAMES = {
    1: "Production & Batch Control", 2: "Procurement & Suppliers", 3: "Bill of Materials",
    4: "Cutting & Planning", 5: "Dyeing & Processing", 6: "Finishing & Treatment",
    7: "Quality Control", 8: "Production Planning", 9: "Maintenance & OEE",
    10: "Warehouse & Rack Mapping", 11: "Inventory & Stock", 12: "Sales & Order Desk",
    13: "CBAM & Carbon Telemetry", 14: "Sustainability & DPP", 15: "Compliance & Audit",
    16: "Finance & Costing", 17: "HR & Workforce", 18: "Dispatch & Logistics",
    19: "Customer Service", 20: "Executive Control Tower",
}


def _user_id():
    return session.get("user_id")


def _data():
    uid = _user_id()
    with db_connection() as c:
        orders = c.execute("SELECT * FROM orders WHERE user_id=? ORDER BY id DESC", (uid,)).fetchall()
        logs = c.execute("SELECT * FROM module_logs WHERE user_id=? ORDER BY id DESC", (uid,)).fetchall()
        cbam = c.execute("SELECT c.*, o.buyer_name, o.fabric_type FROM cbam_logs c JOIN orders o ON o.id=c.order_id WHERE o.user_id=? ORDER BY c.id DESC", (uid,)).fetchall()
        user = c.execute("SELECT company_name, company_logo_path, email FROM users WHERE id=?", (uid,)).fetchone()
    return user, orders, logs, cbam


def _filtered_orders(orders):
    batch_id = request.args.get("batch_id", "").strip()
    buyer = request.args.get("buyer", "").strip().lower()
    status = request.args.get("status", "").strip().upper()
    result = []
    for row in orders:
        if batch_id and str(row["id"]) != batch_id:
            continue
        if buyer and buyer not in row["buyer_name"].lower():
            continue
        if status and row["order_status"].upper() != status:
            continue
        result.append(row)
    return result


@reports_bp.route("/", methods=["GET"])
def center():
    if not _user_id():
        return redirect(url_for("routes.login"))
    user, orders, logs, cbam = _data()
    if not user:
        session.clear()
        flash("Your account session is no longer valid. Please sign in again.", "warning")
        return redirect(url_for("routes.login"))
    return render_template("reports.html", company=user["company_name"], orders=orders, statuses=sorted({r["order_status"] for r in orders}))



def _valid_csrf():
    supplied = request.form.get("csrf_token") or request.headers.get("X-CSRF-Token")
    expected = session.get("csrf_token")
    return bool(supplied and expected and secrets.compare_digest(supplied, expected))


def _row_value(row, *keys):
    normalized = {str(k).strip().lower().replace(" ", "_"): row[k] for k in row.keys()}
    for key in keys:
        value = normalized.get(key)
        if value is not None and str(value).strip() != "":
            return str(value).strip()
    return ""


def _import_rows(file_storage):
    filename = (file_storage.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx")):
        raise ValueError("Import accepts CSV or XLSX files only.")
    raw = file_storage.read()
    if not raw:
        raise ValueError("The selected import file is empty.")
    if filename.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        return list(csv.DictReader(io.StringIO(text)))
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(v).strip() if v is not None else "" for v in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(headers, values)) for values in rows if any(v not in (None, "") for v in values)]


@reports_bp.route("/import", methods=["POST"])
def import_file():
    uid = _user_id()
    if not uid:
        return redirect(url_for("routes.login"))
    if not _valid_csrf():
        flash("Security token validation failed. Refresh the page and retry.", "error")
        return redirect(url_for("reports.center"))
    file_storage = request.files.get("import_file")
    if not file_storage or not file_storage.filename:
        flash("Choose a CSV or XLSX file to import.", "error")
        return redirect(url_for("reports.center"))
    try:
        rows = _import_rows(file_storage)
        if not rows:
            raise ValueError("The import file contains no data rows.")
        imported = 0
        skipped = 0
        with db_connection() as c:
            for row in rows[:1000]:
                buyer = _row_value(row, "buyer_name", "buyer", "customer")
                fabric = _row_value(row, "fabric_type", "fabric", "product")
                quantity_raw = _row_value(row, "quantity_kg", "quantity", "kg")
                if not buyer or not fabric or not quantity_raw:
                    skipped += 1
                    continue
                try:
                    quantity = float(quantity_raw)
                    if quantity <= 0:
                        raise ValueError
                    bom_cost = float(_row_value(row, "total_bom_cost", "bom_cost", "cost") or 0)
                    if bom_cost < 0:
                        raise ValueError
                except (ValueError, TypeError):
                    skipped += 1
                    continue
                qc = (_row_value(row, "qc_grade", "qc") or "PENDING").upper()
                status = (_row_value(row, "order_status", "status") or "ACTIVE").upper()
                if qc not in {"A", "B", "C", "REJECTED", "PENDING"}:
                    qc = "PENDING"
                if status not in {"ACTIVE", "IN PRODUCTION", "READY", "HOLD", "CANCELLED"}:
                    status = "ACTIVE"
                rack = _row_value(row, "warehouse_rack", "rack")
                whrow = _row_value(row, "warehouse_row", "row")
                bin_code = _row_value(row, "warehouse_bin", "bin")
                seed = f"{uid}|{buyer}|{fabric}|{quantity}|{datetime.utcnow().timestamp()}|{secrets.token_hex(8)}"
                import hashlib
                dpp_hash = hashlib.sha256(seed.encode("utf-8")).hexdigest()
                c.execute("INSERT INTO orders (user_id,buyer_name,fabric_type,quantity_kg,dpp_hash,order_status,qc_grade,warehouse_rack,warehouse_row,warehouse_bin,total_bom_cost) VALUES (?,?,?,?,?,?,?,?,?,?,?)", (uid,buyer,fabric,quantity,dpp_hash,status,qc,rack,whrow,bin_code,bom_cost))
                imported += 1
        flash(f"Import complete: {imported} batch record(s) imported; {skipped} row(s) skipped.", "success")
    except Exception as exc:
        flash(f"Import failed: {exc}", "error")
    return redirect(url_for("reports.center"))


def _report_title():
    return request.args.get("title", "Enterprise Operations Report").strip()[:120] or "Enterprise Operations Report"


@reports_bp.route("/pdf")
def pdf_report():
    if not _user_id():
        return redirect(url_for("routes.login"))
    user, orders, logs, cbam = _data()
    if not user:
        session.clear()
        flash("Your account session is no longer valid. Please sign in again.", "warning")
        return redirect(url_for("routes.login"))
    orders = _filtered_orders(orders)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=16*mm, bottomMargin=16*mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SmallGray", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#64748b"), leading=11))
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontSize=20, leading=24, spaceAfter=5))
    story = [Paragraph(user["company_name"], styles["ReportTitle"]), Paragraph(_report_title(), styles["Heading2"]), Paragraph(f"Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}", styles["SmallGray"]), Spacer(1, 8*mm)]

    summary = [["Batches", str(len(orders))], ["Total Quantity (KG)", f"{sum(float(r['quantity_kg']) for r in orders):,.2f}"], ["Active Batches", str(sum(1 for r in orders if r['order_status']=='ACTIVE'))], ["CBAM Records", str(len(cbam))]]
    t = Table(summary, colWidths=[55*mm, 45*mm])
    t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#f1f5f9")), ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#cbd5e1")), ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"), ("FONTNAME", (1,0), (1,-1), "Helvetica"), ("PADDING", (0,0), (-1,-1), 6)]))
    story += [t, Spacer(1, 8*mm), Paragraph("Production Batch Register", styles["Heading2"])]
    data = [["Batch", "Buyer", "Fabric", "KG", "QC", "Status", "Warehouse", "BOM Cost"]]
    for r in orders:
        wh = " / ".join([r["warehouse_rack"] or "—", r["warehouse_row"] or "—", r["warehouse_bin"] or "—"])
        data.append([f"BATCH-{r['id']}", r["buyer_name"], r["fabric_type"], f"{float(r['quantity_kg']):,.2f}", r["qc_grade"], r["order_status"], wh, f"${float(r['total_bom_cost']):,.2f}"])
    if len(data) == 1:
        data.append(["No matching records", "", "", "", "", "", "", ""])
    table = Table(data, repeatRows=1, colWidths=[20*mm, 28*mm, 38*mm, 17*mm, 15*mm, 24*mm, 30*mm, 22*mm])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f172a")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 6.5), ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#cbd5e1")), ("VALIGN", (0,0), (-1,-1), "TOP"), ("PADDING", (0,0), (-1,-1), 4)]))
    story.append(table)
    story += [Spacer(1, 8*mm), Paragraph("Operational Telemetry", styles["Heading2"])]
    logdata = [["Module", "Reference", "Value", "Notes", "Time"]]
    for r in logs[:100]:
        logdata.append([MODULE_NAMES.get(r["module_id"], f"Module {r['module_id']}"), r["reference"], r["value"], r["notes"] or "", r["created_at"]])
    if len(logdata) == 1:
        logdata.append(["No records", "", "", "", ""])
    lt = Table(logdata, repeatRows=1, colWidths=[30*mm, 30*mm, 42*mm, 55*mm, 25*mm])
    lt.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f172a")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTSIZE", (0,0), (-1,-1), 6), ("GRID", (0,0), (-1,-1), .3, colors.HexColor("#cbd5e1")), ("VALIGN", (0,0), (-1,-1), "TOP"), ("PADDING", (0,0), (-1,-1), 3)]))
    story.append(lt)
    doc.build(story)
    buffer.seek(0)
    return send_file(buffer, mimetype="application/pdf", as_attachment=True, download_name=f"EcoTextile_Report_{datetime.now():%Y%m%d_%H%M}.pdf")


@reports_bp.route("/excel")
def excel_report():
    if not _user_id():
        return redirect(url_for("routes.login"))
    user, orders, logs, cbam = _data()
    if not user:
        session.clear()
        flash("Your account session is no longer valid. Please sign in again.", "warning")
        return redirect(url_for("routes.login"))
    orders = _filtered_orders(orders)
    wb = Workbook()
    ws = wb.active
    ws.title = "Production Batches"
    headers = ["Batch ID", "Buyer", "Fabric Type", "Quantity KG", "DPP Hash", "QC Grade", "Status", "Rack", "Row", "Bin", "BOM Cost"]
    ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="0F172A"); c.font = Font(color="FFFFFF", bold=True)
    for r in orders:
        ws.append([f"BATCH-{r['id']}", r["buyer_name"], r["fabric_type"], float(r["quantity_kg"]), r["dpp_hash"], r["qc_grade"], r["order_status"], r["warehouse_rack"] or "", r["warehouse_row"] or "", r["warehouse_bin"] or "", float(r["total_bom_cost"])])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(max(len(str(c.value or "")) for c in col)+2, 12), 35)
    ws.freeze_panes = "A2"

    wl = wb.create_sheet("Module Logs")
    wl.append(["Module", "Reference", "Value", "Notes", "Created At"])
    for r in logs:
        wl.append([MODULE_NAMES.get(r["module_id"], f"Module {r['module_id']}"), r["reference"], r["value"], r["notes"] or "", r["created_at"]])
    wc = wb.create_sheet("CBAM")
    wc.append(["Batch ID", "Buyer", "Fabric Type", "Emissions Metric", "Target Tax USD"])
    for r in cbam:
        wc.append([f"BATCH-{r['order_id']}", r["buyer_name"], r["fabric_type"], float(r["emissions_metric"]), float(r["target_tax_usd"])])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]: cell.font = Font(bold=True)
        for row in sheet.iter_rows():
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in col)+2, 12), 40)
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"EcoTextile_Export_{datetime.now():%Y%m%d_%H%M}.xlsx")
